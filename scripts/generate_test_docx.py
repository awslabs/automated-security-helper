# Copyright Amazon.com, Inc. or its affiliates. All Rights Reserved.
# SPDX-License-Identifier: Apache-2.0

"""Generate synthetic test data docx for ferret-scan plugin testing.

This script creates a docx file with synthetic PII-like data covering
all ferret-scan check categories. All data is entirely fictional and
generated programmatically — no real PII is used.

Usage:
    uv run python scripts/generate_test_docx.py
"""

import datetime
import io
import random  # nosec B311 — seeded PRNG for reproducible synthetic test data
import string
import zipfile
from pathlib import Path

from docx import Document

# Anchored on this file, not on the working directory: a relative path wrote the
# fixture into whatever directory the script was run from, silently creating a
# new tests/test_data/... tree there when that was not the repository root.
REPO_ROOT = Path(__file__).resolve().parents[1]
OUTPUT_PATH = (
    REPO_ROOT
    / "tests"
    / "test_data"
    / "scanners"
    / "ferret-scan"
    / "synthetic-pii-test.docx"
)

#: The PRNG seed. Applied at the start of main(), not at import: seeded once at
#: import, a second main() in one process drew different data, and importing
#: the module reseeded the process-global PRNG for whoever else used it.
SEED = 42

#: The timestamps written into the document, so two runs save the same bytes.
#: The zip date_time tuple is the earliest a zip entry can carry.
FIXED_ZIP_DATE_TIME = (1980, 1, 1, 0, 0, 0)
FIXED_DOCUMENT_TIME = datetime.datetime(2024, 1, 1, 0, 0, 0)

# ============================================================================
# Synthetic data generators
# ============================================================================

FIRST_NAMES = [
    "Alice",
    "Bob",
    "Carol",
    "David",
    "Eve",
    "Frank",
    "Grace",
    "Henry",
    "Iris",
    "Jack",
    "Karen",
    "Leo",
    "Mia",
    "Noah",
    "Olivia",
    "Paul",
    "Quinn",
    "Rosa",
    "Sam",
    "Tina",
    "Uma",
    "Victor",
    "Wendy",
    "Xavier",
    "Yara",
    "Zane",
]

LAST_NAMES = [
    "Anderson",
    "Brown",
    "Clark",
    "Davis",
    "Evans",
    "Foster",
    "Garcia",
    "Harris",
    "Ingram",
    "Johnson",
    "King",
    "Lopez",
    "Miller",
    "Nelson",
    "Ortiz",
    "Parker",
    "Quinn",
    "Roberts",
    "Smith",
    "Taylor",
    "Underwood",
    "Vance",
    "Williams",
    "Xu",
    "Young",
    "Zhang",
]


def fake_name():
    return f"{random.choice(FIRST_NAMES)} {random.choice(LAST_NAMES)}"  # nosec B311


def fake_ssn():
    """Generate a fake SSN in the shape sample.txt uses (123-45-6789).

    The area number is drawn from the range a format validator accepts: 900-999
    is the range the SSA never issues, and a detector that applies that rule
    (ferret-scan's does) scored every fixture SSN below detection, so the SSN
    section exercised nothing. 000 and 666 are likewise never issued and are
    skipped. The numbers are still fabricated: no real SSN is looked up or used.
    """
    area = 666
    while area == 666:
        area = random.randint(100, 899)  # nosec B311
    group = random.randint(10, 99)  # nosec B311
    serial = random.randint(1000, 9999)  # nosec B311
    return f"{area}-{group}-{serial}"


def luhn_check_digit(partial: str) -> str:
    """The digit that makes ``partial`` + digit pass the Luhn check.

    ferret-scan validates every card-number candidate with Luhn before it
    reports one, so a fixture number with a random last digit is dropped nine
    times in ten and the fixture under-exercises the detector it exists to test.
    """
    total = 0
    for index, char in enumerate(reversed(partial)):
        digit = int(char)
        if index % 2 == 0:
            digit *= 2
            if digit > 9:
                digit -= 9
        total += digit
    return str((10 - total % 10) % 10)


def luhn_valid(number: str) -> bool:
    """Whether the digits of ``number`` (separators ignored) pass the Luhn check."""
    digits = "".join(char for char in number if char.isdigit())
    return len(digits) > 1 and luhn_check_digit(digits[:-1]) == digits[-1]


def _random_digits(count: int) -> str:
    return "".join(str(random.randint(0, 9)) for _ in range(count))  # nosec B311


def fake_visa():
    """Generate a fake Visa-like number (starts with 4, Luhn-valid)."""
    partial = "4" + _random_digits(14)
    digits = partial + luhn_check_digit(partial)
    return f"{digits[0:4]}-{digits[4:8]}-{digits[8:12]}-{digits[12:16]}"


def fake_mastercard():
    """Generate a fake MasterCard-like number (starts with 5, Luhn-valid)."""
    partial = "5" + str(random.randint(1, 5)) + _random_digits(13)  # nosec B311
    digits = partial + luhn_check_digit(partial)
    return f"{digits[0:4]}-{digits[4:8]}-{digits[8:12]}-{digits[12:16]}"


def fake_amex():
    """Generate a fake AMEX-like number (starts with 37, Luhn-valid)."""
    partial = "37" + _random_digits(12)
    digits = partial + luhn_check_digit(partial)
    return f"{digits[0:4]}-{digits[4:10]}-{digits[10:15]}"


def fake_discover():
    """Generate a fake Discover-like number (starts with 6011, Luhn-valid)."""
    partial = "6011" + _random_digits(11)
    digits = partial + luhn_check_digit(partial)
    return f"{digits[0:4]} {digits[4:8]} {digits[8:12]} {digits[12:16]}"


def fake_email(name):
    first, last = name.lower().split()
    domains = ["example.com", "test.org", "sample.net", "demo.io"]
    return f"{first}.{last}@{random.choice(domains)}"  # nosec B311


def fake_phone():
    area = random.randint(200, 999)  # nosec B311
    prefix = random.randint(200, 999)  # nosec B311
    line = random.randint(1000, 9999)  # nosec B311
    return f"({area}) {prefix}-{line}"


def fake_passport():
    letter = random.choice(string.ascii_uppercase)  # nosec B311
    digits = "".join(str(random.randint(0, 9)) for _ in range(8))  # nosec B311
    return f"{letter}{digits}"


#: A TD3 (passport) machine-readable zone line is exactly this many characters.
TD3_LINE_LENGTH = 44

#: The two MRZ-style passport lines, each padded with the `<` filler to the TD3
#: length. Fictional holders.
MRZ_LINES = (
    "P<USASMITH<<JOHN".ljust(TD3_LINE_LENGTH, "<"),
    "P<GBRJONES<<SARAH".ljust(TD3_LINE_LENGTH, "<"),
)


def _fixed_zip_timestamps(docx_bytes: bytes) -> bytes:
    """``docx_bytes`` rewritten so every zip member carries one fixed timestamp.

    python-docx saves the package through ``zipfile.writestr`` with a bare name,
    which stamps each member with the current time, so a seeded run still wrote
    different bytes on every save and the committed fixture churned on each
    regeneration. The members, their order and their contents are unchanged.
    """
    rewritten = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(docx_bytes)) as source:
        with zipfile.ZipFile(rewritten, "w") as target:
            for member in source.infolist():
                info = zipfile.ZipInfo(member.filename, date_time=FIXED_ZIP_DATE_TIME)
                info.compress_type = member.compress_type
                info.external_attr = member.external_attr
                target.writestr(info, source.read(member))
    return rewritten.getvalue()


def write_reproducible_docx(doc, output_path: Path) -> None:
    """Save ``doc`` to ``output_path`` with byte-reproducible metadata.

    The core properties python-docx stamps at save time (created, modified,
    last_modified_by, revision) are pinned, and the zip member timestamps are
    rewritten to a fixed value, so two runs with the same seed write the same
    bytes.
    """
    core = doc.core_properties
    core.created = FIXED_DOCUMENT_TIME
    core.modified = FIXED_DOCUMENT_TIME
    core.last_printed = FIXED_DOCUMENT_TIME
    core.last_modified_by = "ASH Test Data Generator"
    core.revision = 1
    buffer = io.BytesIO()
    doc.save(buffer)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_bytes(_fixed_zip_timestamps(buffer.getvalue()))


def main(output_path: Path = OUTPUT_PATH):
    # Seeded here so every run of main() draws the same sequence, whatever ran
    # before it in the process.
    random.seed(SEED)
    doc = Document()

    # Clear default metadata
    core = doc.core_properties
    core.author = "ASH Test Data Generator"
    core.title = "Synthetic PII Test Data for Ferret-Scan"
    core.subject = "Test data - all PII is fictional"
    core.comments = (
        "Generated by automated_security_helper for plugin testing. "
        "All data is synthetic. No real PII. Apache-2.0 licensed."
    )

    # Title
    doc.add_heading("Synthetic PII Test Data", level=1)
    doc.add_paragraph(
        "This document contains synthetic (fake) PII data for testing "
        "ferret-scan's detection capabilities. All data is entirely fictional "
        "and generated programmatically. No real personal information is included."
    )
    doc.add_paragraph(
        "Copyright Amazon.com, Inc. or its affiliates. All Rights Reserved. "
        "SPDX-License-Identifier: Apache-2.0"
    )

    # ---- Section 1: Plain text PII that matches ferret-scan patterns ----
    # These use the same formats that work in sample.txt
    doc.add_heading("Section 1: Credit Card Numbers", level=2)

    # Generate card numbers programmatically to avoid hardcoded strings
    # that trigger secret detectors in CI (which runs --ignore-suppressions)
    for _ in range(5):
        doc.add_paragraph(f"Visa: {fake_visa()}")

    for _ in range(5):
        doc.add_paragraph(f"MasterCard: {fake_mastercard()}")

    for _ in range(3):
        doc.add_paragraph(f"American Express: {fake_amex()}")

    for _ in range(2):
        doc.add_paragraph(f"Discover: {fake_discover()}")

    doc.add_heading("Section 2: Social Security Numbers", level=2)
    for _ in range(5):
        doc.add_paragraph(f"SSN: {fake_ssn()}")

    doc.add_heading("Section 3: Passport Numbers", level=2)
    for _ in range(5):
        doc.add_paragraph(f"Passport: {fake_passport()}")
    # MRZ-style lines, padded to the 44 characters a TD3 line has. Written out
    # by hand they were 45, one filler too many, which a length-checking MRZ
    # detector rejects.
    for mrz_line in MRZ_LINES:
        doc.add_paragraph(f"Passport: {mrz_line}")

    doc.add_heading("Section 4: Phone Numbers", level=2)
    for _ in range(5):
        doc.add_paragraph(f"Phone: {fake_phone()}")

    doc.add_heading("Section 5: Contact Records", level=2)

    # The narrative that quotes a card's last four digits is built from the card
    # it embeds, so the two agree, as the mirrored sample.txt line does. A fixed
    # "ending in 3020" beside a random number contradicted itself.
    approved_card = fake_mastercard()
    approved_last_four = approved_card[-4:]

    # Free-text narratives with embedded PII (generated to avoid hardcoded secrets)
    narratives = [
        (
            f"Customer {fake_name()} (SSN: {fake_ssn()}) reported unauthorized charges "
            f"on Visa card {fake_visa()}. Contact at {fake_phone()}."
        ),
        (
            f"Employee {fake_name()} submitted expense report. Corporate MasterCard "
            f"{fake_mastercard()} used for travel. Phone: {fake_phone()}."
        ),
        (
            f"Account holder {fake_name()} requested replacement AMEX card {fake_amex()}. "
            f"Verified identity with SSN {fake_ssn()} and passport {fake_passport()}."
        ),
        (
            f"HR record for {fake_name()}: SSN {fake_ssn()}, phone {fake_phone()}. "
            f"Travel document: passport {fake_passport()}. Company Visa: {fake_visa()}."
        ),
        (
            f"Please charge Visa card {fake_visa()} for the full amount. "
            f"The MasterCard ending in {approved_last_four} ({approved_card}) has "
            f"been approved."
        ),
        (
            f"AMEX: {fake_amex()} - Expiration: 12/25. "
            f"Discover card {fake_discover()} authorized for recurring payments."
        ),
    ]
    for n in narratives:
        doc.add_paragraph(n)

    # ---- Section 6: Bulk records to add volume ----
    doc.add_heading("Section 6: Extended Records", level=2)

    NUM_RECORDS = 200

    for _ in range(NUM_RECORDS):
        name = fake_name()
        doc.add_paragraph(
            f"{name} | {fake_ssn()} | {fake_visa()} | {fake_mastercard()} | "
            f"{fake_amex()} | {fake_email(name)} | {fake_phone()} | {fake_passport()}"
        )

    # Save, byte-reproducibly: the same seed writes the same file.
    write_reproducible_docx(doc, output_path)
    size_kb = output_path.stat().st_size / 1024
    print(f"Generated: {output_path} ({size_kb:.0f} KB)")


if __name__ == "__main__":
    main()
